#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the ODI-branded launch webinar deck from the ODI PPTX template.

This script intentionally edits the retained template slides in-place:
it opens the ODI corporate template, keeps/reorders the 18 selected slides by
editing presentation.xml's <p:sldIdLst>, then replaces text and pictures on
those native slides. It prints one progress line per slide plus non-fatal
warnings for template-shape mismatches.
"""

from __future__ import annotations

import copy
import re
import struct
import sys
import traceback
from pathlib import Path
from typing import Iterable, Sequence


REPO_ROOT = Path(__file__).resolve().parent

TEMPLATE_PATH = REPO_ROOT / "pptx-build" / "odi-template.pptx"
BW_DIR = REPO_ROOT / "pptx-build" / "bw"
IMG_DIR = REPO_ROOT / "public" / "webinar" / "img"
OUTPUT_PATH = REPO_ROOT / "webinar-odi.pptx"

ODI_BLUE_HEX = "072589"
FIGMA_URL = (
    "https://www.figma.com/proto/ZYp8o7zamQX9EGptcCBdw1/"
    "Solid-interfaces--ODI-?page-id=19732%3A27511&node-id=19732-27512&"
    "viewport=-6417%2C1588%2C0.26&t=dPgT5VOyEOAnJKHo-1&scaling=scale-down&"
    "content-scaling=fixed&starting-point-node-id=19732%3A52257"
)

# ---------------------------------------------------------------------------
# EDITABLE SLIDE COPY
#
# Non-programmers should be able to update deck wording here and rerun the
# builder. This copy mirrors the approved deployed deck at
# public/webinar/index.html.
# ---------------------------------------------------------------------------
CONTENT = {
    "s01": {
        "title": "One copy of the truth",
        "strap": "Sir Tim Berners-Lee on Solid and the future of enterprise data",
        "meta": [
            "EDM Association + the Open Data Institute",
            "14 July 2026 · 10:00 am EST / 3:00 pm UK",
            "This session will not be recorded",
        ],
    },
    "s02": {
        "title": "Run of show",
        "agenda": [
            {
                "number": "1.",
                "title": "Welcome & introductions",
                "speaker": "Emma Thwaites · ODI",
                "duration": "5 min",
            },
            {
                "number": "2.",
                "title": "Introduction to Solid",
                "speaker": "Jesse Wright · ODI",
                "duration": "10 min",
            },
            {
                "number": "3.",
                "title": "In conversation: Solid & the future of enterprise data",
                "speaker": "Sir Tim Berners-Lee & John Bruce · Inrupt",
                "duration": "20 min",
            },
            {
                "number": "4.",
                "title": "Audience Q&A",
                "speaker": "Hosted by Emma Thwaites",
                "duration": "10 min",
            },
            {
                "number": "5.",
                "title": "The Solid Community of Practice",
                "speaker": "Jim Halcomb & Elisa Kendall · EDM Association",
                "duration": "5 min",
            },
        ],
        "footer": "14 July 2026 · 10:00 am EST / 3:00 pm UK · not recorded · put your questions in the chat throughout",
    },
    "s03": {
        "kicker": "Emma Thwaites · welcome & introductions",
        "speakers": [
            {
                "image": "emma-thwaites.jpg",
                "name": "Emma Thwaites",
                "role": "Moderator · Strategic Advisor, Global Policy & Corporate Affairs, ODI",
            },
            {
                "image": "tim-berners-lee.jpg",
                "name": "Sir Tim Berners-Lee",
                "role": "Inventor of the World Wide Web · Co-founder & CTO, Inrupt Inc.",
            },
            {
                "image": "john-bruce.jpg",
                "name": "John Bruce",
                "role": "Co-founder & CEO, Inrupt Inc.",
            },
            {
                "image": "jesse-wright.jpg",
                "name": "Jesse Wright",
                "role": "Project Lead – Solid, ODI",
            },
            {
                "image": "jim-halcomb.jpg",
                "name": "Jim Halcomb",
                "role": "Chief Research & Development Officer, EDM Association",
            },
            {
                "image": "elisa-kendall.jpg",
                "name": "Elisa Kendall",
                "role": "Senior Ontologist, EDM Association",
            },
        ],
    },
    "s04": {
        "kicker": "Emma Thwaites · a brief word on the ODI",
        "title": "Steward of Solid since 2024",
        "body": [
            "The Open Data Institute is a mission-driven non-profit, co-founded in 2012 by Sir Tim Berners-Lee and Sir Nigel Shadbolt, working towards an open, trustworthy data ecosystem where data works for everyone.",
            "Since October 2024, the ODI has stewarded Solid: its standards at the W3C, its codebases, and its global community. Its stewardship aims to establish ethical data governance, addressing privacy, transparency and control so data can be used responsibly for individuals, society and the economy.",
        ],
        "trusted_label": "Trusted by clients and partners",
    },
    "s05": {},
    "s06": {
        "title": "Introduction to Solid",
        "subtitle": "What it is, the problem it solves, and a live demo",
        "who": "Jesse Wright · Project Lead – Solid, the ODI · 10 minutes",
    },
    "s07": {
        "title": "What is Solid",
        "body": [
            "Solid is an open-source **project**, open **standard**, and **movement** for a **user-centric data** storage system, originating from Sir Tim Berners-Lee's lab at MIT CSAIL.",
            "Your data lives in a **Pod**: a personal or organisational datastore with fine-grained access controls. **AI agents and apps** read and write only **with permission**, **breaking down data silos** and keeping data under the control of whoever it describes.",
            "The ODI stewards Solid's standards, codebases, and communities.",
        ],
        "figure_caption": "Solid extends the existing web architecture",
    },
    "s08": {
        "title": "Every service keeps its own copy of your data",
        "lead": "The same data is collected, copied and secured many times over, yet is rarely current, portable or under the control of the person or organisation it describes. The result:",
        "bullets": [
            "Silos and duplication",
            "More places for data to be breached",
            "Vendor lock-in",
            "All-or-nothing consent",
            "Having to prove your identity from scratch at every institution",
        ],
        "fix": "**Solid fixes this by decoupling data from applications.** Services read what they're allowed to use from the person's or organisation's Pod, instead of keeping another copy.",
        "ai_line": "As AI agents act for us, they need current data the person can still control — Solid gives them a way to use it with permission.",
        "figure_caption": "Walled Gardens, by David Simonds",
    },
    "s09": {
        "flag": "Live demo",
        "title": "Solid in action — live walkthrough",
        "footer": "Driven & narrated live by Jesse Wright",
    },
    "s10": {
        "kicker": "In conversation · 20 minutes",
        "title": "Solid & the future of enterprise data",
        "footer": "On the table: what Solid is and why it was built · early deployments across government and healthcare · why it matters even more in the age of AI. Put your questions in the chat — we'll come to them next.",
        "people": [
            {
                "image": "tim-berners-lee.jpg",
                "name": "Sir Tim Berners-Lee",
                "role": "Inventor of the World Wide Web · Co-founder & CTO, Inrupt Inc.",
            },
            {"image": "john-bruce.jpg", "name": "John Bruce", "role": "Co-founder & CEO, Inrupt Inc."},
        ],
    },
    "s11": {
        "kicker": "Audience Q&A · 10 minutes",
        "title": "Your questions",
        "body": [
            "Hosted by **Emma Thwaites**.",
            "Drop your questions in the chat — the moderator will bring them to the panel.",
        ],
        "footer": "On the panel: Sir Tim Berners-Lee · John Bruce",
    },
    "s12": {
        "title": "The Solid Community of Practice",
        "subtitle": "What it is, what we'll build, and how to take part",
        "who": "Jim Halcomb & Elisa Kendall · EDM Association · 5 minutes",
    },
    "s13": {
        "kicker": "Jim Halcomb & Elisa Kendall",
        "title": "Putting one copy of the truth into practice",
        "lead": "A **community of practice** is a neutral, vendor-independent group where finance, healthcare and government work alongside the people building Solid to make it production-ready for regulated data. Built on 300+ participants in the W3C Solid Community Group; every output MIT-licensed.",
        "cards": [
            {
                "label": "Real pilots",
                "title": "Real pilots per sector",
                "body": "Demonstrable deployments in finance, healthcare and government — not slideware.",
            },
            {
                "label": "Standards",
                "title": "Aligned to sector standards",
                "body": "How Solid connects to the models these industries already rely on — starting with the Financial Industry Business Ontology (FIBO) in finance.",
            },
            {
                "label": "Certify",
                "title": "Independent certification",
                "body": "A trusted, neutral way to certify that a Solid service meets a sector's requirements.",
            },
        ],
    },
    "s14": {
        "kicker": "Jim Halcomb & Elisa Kendall · focus A",
        "title": "Real pilots, in each sector",
        "points": [
            {
                "heading": "Real deployments",
                "body": "End-to-end pilots that can be shown, tested and written up as reference implementations, with named champions, case studies and reports, and webinars so others can follow.",
            },
            {
                "heading": "Finance first",
                "body": "Customer onboarding and lending decisions: data checked against the customer's own Pod, reusable identity and eligibility checks, and credentials issued by a bank that the customer can use elsewhere.",
            },
            {
                "heading": "Then healthcare and government",
                "body": "Apply the same pattern where repeated checks and copied records create cost and risk, starting with practical services that can be demonstrated.",
            },
        ],
    },
    "s15": {
        "kicker": "Jim Halcomb & Elisa Kendall · focus B",
        "title": "Aligned to the standards each sector already trusts",
        "points": [
            {
                "heading": "Connect Solid to trusted sector models",
                "body": "Map Solid to the data models each sector already uses, rather than asking organisations to adopt a new vocabulary.",
            },
            {
                "heading": "Finance starts with the Financial Industry Business Ontology (FIBO)",
                "body": "Use FIBO first, so data in a Pod is understood the same way by every participating financial institution.",
            },
            {
                "heading": "Describe what a Pod holds",
                "body": "Let a Pod describe the data products it contains and how access should be requested. Healthcare and government follow.",
            },
        ],
    },
    "s16": {
        "kicker": "Jim Halcomb & Elisa Kendall · focus C",
        "title": "A trusted, independent certification",
        "points": [
            {
                "heading": "Solid-CDMC guidance and certification",
                "body": "Adapt the EDM Council's Cloud Data Management Capabilities (CDMC) controls to Solid services, so Pod providers and application teams know what evidence is expected for governance, access control and operational readiness. A Solid-CDMC certification would confirm that a Solid service meets those adapted controls.",
            },
            {
                "heading": "A neutral mark",
                "body": "Set up a model to show that Solid services are fit for deployment in a particular sector, including security, conformance and sector-specific requirements.",
            },
            {
                "heading": "Re-use a proven partner model",
                "body": "Run assessment through EDMA's established authorised-partner model: trained partners assess Solid services, with a neutral body overseeing certification across sectors.",
            },
        ],
    },
    "s17": {
        "kicker": "Jim Halcomb · get involved",
        "title": "Lead how Solid reshapes your sector",
        "badge": "EDMConnect",
        "register": "Register via EDMConnect",
        "url": "hub.edmcouncil.org",
        "body": "Registration is open to everyone; EDM Association membership is not required. Once registered, join the community calls and bring a pilot, a use case for connecting Solid to sector standards, or a certification use case when ready.",
        "cochairs": "Co-chairs: Jim Halcomb & Elisa Kendall (EDM Association) · Roberto Breitman, Stuart Coleman & Jesse Wright (ODI) · Jacobus Geluk (agnos.ai).",
    },
    "s18": {
        "title": "Join the Solid Community of Practice",
        "title_lines": ["Join the Solid", "Community of Practice"],
        "subheader": "Register: hub.edmcouncil.org · Read more: edma-microsite.vercel.app · Co-convened by the EDM Association and the ODI · not recorded · MIT-licensed",
    },
}

# Our webinar slide number -> 1-based template slide index.
TEMPLATE_INDICES = [3, 4, 59, 46, 26, 15, 32, 9, 53, 57, 18, 16, 41, 6, 7, 8, 20, 14]


def die(message: str, code: int = 1) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(code)


def import_pptx():
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
        from pptx.enum.text import PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        try:
            from pptx.enum.text import MSO_AUTO_SIZE
        except ImportError:  # pragma: no cover - version compatibility
            MSO_AUTO_SIZE = None
    except Exception as exc:
        print("ERROR: python-pptx is required to build this deck.", file=sys.stderr)
        print("Install it in a virtualenv, for example: python -m pip install python-pptx", file=sys.stderr)
        print(f"Import failure: {exc!r}", file=sys.stderr)
        raise

    return {
        "Presentation": Presentation,
        "MSO_SHAPE": MSO_SHAPE,
        "MSO_SHAPE_TYPE": MSO_SHAPE_TYPE,
        "PP_ALIGN": PP_ALIGN,
        "MSO_AUTO_SIZE": MSO_AUTO_SIZE,
        "qn": qn,
        "Inches": Inches,
        "Pt": Pt,
        "RGBColor": RGBColor,
    }


def validate_inputs() -> str:
    required = [
        TEMPLATE_PATH,
        BW_DIR / "emma-thwaites.jpg",
        BW_DIR / "tim-berners-lee.jpg",
        BW_DIR / "john-bruce.jpg",
        BW_DIR / "jesse-wright.jpg",
        BW_DIR / "jim-halcomb.jpg",
        BW_DIR / "elisa-kendall.jpg",
        IMG_DIR / "clients.png",
        IMG_DIR / "data-spectrum.jpg",
        IMG_DIR / "walled-gardens.png",
        IMG_DIR / "solid-pods.png",
        IMG_DIR / "solid-demo.jpg",
    ]

    missing = [path for path in required if not path.exists()]
    if missing:
        for path in missing:
            print(f"[missing] {path}", file=sys.stderr)
        die("One or more required input files are missing.")

    return ""


def as_lines(value: str | Sequence[str]) -> list[str]:
    if isinstance(value, str):
        return [value]
    return list(value)


def normalize_text(value: str) -> str:
    """Normalise template text so matching survives run splits and line breaks."""
    return re.sub(r"\s+", " ", value.replace("\xa0", " ")).strip().casefold()


def parse_bold_markers(text: str) -> list[tuple[str, bool | None]]:
    """Parse a tiny subset of Markdown: **bold** markers only."""
    segments: list[tuple[str, bool | None]] = []
    cursor = 0
    bold = False
    while cursor < len(text):
        marker = text.find("**", cursor)
        if marker == -1:
            if cursor < len(text):
                segments.append((text[cursor:], True if bold else None))
            break
        if marker > cursor:
            segments.append((text[cursor:marker], True if bold else None))
        bold = not bold
        cursor = marker + 2
    return [(segment, flag) for segment, flag in segments if segment]


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs)


def walk_shapes(shapes) -> Iterable:
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from walk_shapes(shape.shapes)


def shape_label(shape) -> str:
    name = getattr(shape, "name", "")
    shape_id = getattr(shape, "shape_id", "?")
    return f"id={shape_id} name={name!r}"


def find_shapes_containing(slide, needle: str, *, exact: bool = False) -> list:
    matches = []
    normalized_needle = normalize_text(needle)
    for shape in walk_shapes(slide.shapes):
        text = shape_text(shape)
        if not text:
            continue
        normalized_text = normalize_text(text)
        matched = normalized_text == normalized_needle if exact else normalized_needle in normalized_text
        if matched:
            matches.append(shape)
    return matches


def delete_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def delete_shapes_containing(slide, needles: Sequence[str]) -> int:
    deleted = 0
    normalized_needles = [normalize_text(needle) for needle in needles]
    for shape in list(walk_shapes(slide.shapes)):
        text = shape_text(shape)
        normalized_text = normalize_text(text)
        if text and any(needle in normalized_text for needle in normalized_needles):
            delete_shape(shape)
            deleted += 1
    return deleted


def capture_text_templates(shape):
    templates = []
    for paragraph in shape.text_frame.paragraphs:
        ppr = copy.deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
        rpr = None
        if paragraph.runs:
            first_run = paragraph.runs[0]
            rpr = copy.deepcopy(first_run._r.rPr) if first_run._r.rPr is not None else None
        templates.append((ppr, rpr))
    return templates or [(None, None)]


def apply_paragraph_template(paragraph, template) -> None:
    ppr, _ = template
    if ppr is None:
        return
    existing = paragraph._p.pPr
    if existing is not None:
        paragraph._p.remove(existing)
    paragraph._p.insert(0, copy.deepcopy(ppr))


def apply_run_template(run, template) -> None:
    _, rpr = template
    if rpr is None:
        return
    existing = run._r.rPr
    if existing is not None:
        run._r.remove(existing)
    run._r.insert(0, copy.deepcopy(rpr))


def set_rich_text(
    shape,
    paragraphs: Sequence[str | Sequence[tuple[str, bool | None]]],
    *,
    style_indices: Sequence[int] | None = None,
) -> None:
    """Replace all text in a shape while preserving paragraph/run styling templates."""
    if not getattr(shape, "has_text_frame", False):
        raise ValueError(f"Shape {shape_label(shape)} has no text frame.")

    templates = capture_text_templates(shape)
    text_frame = shape.text_frame
    text_frame.clear()

    for idx, paragraph_data in enumerate(paragraphs):
        paragraph = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
        template_index = style_indices[idx] if style_indices and idx < len(style_indices) else idx
        template = templates[min(template_index, len(templates) - 1)]
        apply_paragraph_template(paragraph, template)

        if isinstance(paragraph_data, str):
            segments = [(paragraph_data, None)]
        else:
            segments = list(paragraph_data)

        # text_frame.clear() leaves no runs in the first paragraph under current python-pptx,
        # but clear any accidental placeholder text defensively.
        for run in list(paragraph.runs):
            run.text = ""

        for segment, bold in segments:
            run = paragraph.add_run()
            apply_run_template(run, template)
            run.text = segment
            if bold is not None:
                run.font.bold = bold


def set_text(shape, value: str | Sequence[str], *, style_indices: Sequence[int] | None = None) -> None:
    set_rich_text(shape, as_lines(value), style_indices=style_indices)


def set_markdown_text(shape, paragraphs: Sequence[str], *, style_indices: Sequence[int] | None = None) -> None:
    set_rich_text(shape, [parse_bold_markers(paragraph) for paragraph in paragraphs], style_indices=style_indices)


def replace_first(slide, needle: str, value: str | Sequence[str], warnings: list[str], slide_name: str) -> None:
    matches = find_shapes_containing(slide, needle)
    if not matches:
        warnings.append(f"{slide_name}: could not find text containing {needle!r}")
        return
    set_text(matches[0], value)


def replace_first_of(
    slide,
    needles: Sequence[str],
    value: str | Sequence[str],
    warnings: list[str],
    slide_name: str,
    role: str,
) -> None:
    for needle in needles:
        matches = find_shapes_containing(slide, needle)
        if matches:
            set_text(matches[0], value)
            return
    warnings.append(f"{slide_name}: could not find {role}; tried {', '.join(repr(needle) for needle in needles)}")


def replace_all_exact(slide, needle: str, value: str | Sequence[str], warnings: list[str], slide_name: str) -> None:
    matches = find_shapes_containing(slide, needle, exact=True)
    if not matches:
        warnings.append(f"{slide_name}: could not find exact text {needle!r}")
        return
    for shape in matches:
        set_text(shape, value)


def sorted_text_shapes(slide, needle: str, *, exact: bool = False) -> list:
    return sorted(find_shapes_containing(slide, needle, exact=exact), key=lambda shape: (shape.left, shape.top))


def picture_shapes(slide, MSO_SHAPE_TYPE) -> list:
    return sorted(
        [shape for shape in walk_shapes(slide.shapes) if shape.shape_type == MSO_SHAPE_TYPE.PICTURE],
        key=lambda shape: (shape.left, shape.top),
    )


def png_size(path: Path) -> tuple[int, int] | None:
    with path.open("rb") as fh:
        header = fh.read(24)
    if header[:8] != b"\x89PNG\r\n\x1a\n":
        return None
    return struct.unpack(">II", header[16:24])


def jpeg_size(path: Path) -> tuple[int, int] | None:
    with path.open("rb") as fh:
        if fh.read(2) != b"\xff\xd8":
            return None
        while True:
            byte = fh.read(1)
            if not byte:
                return None
            if byte != b"\xff":
                continue
            marker = fh.read(1)
            while marker == b"\xff":
                marker = fh.read(1)
            if marker in [b"\xd8", b"\xd9"]:
                continue
            length_bytes = fh.read(2)
            if len(length_bytes) != 2:
                return None
            length = struct.unpack(">H", length_bytes)[0]
            if marker in [
                b"\xc0",
                b"\xc1",
                b"\xc2",
                b"\xc3",
                b"\xc5",
                b"\xc6",
                b"\xc7",
                b"\xc9",
                b"\xca",
                b"\xcb",
                b"\xcd",
                b"\xce",
                b"\xcf",
            ]:
                data = fh.read(length - 2)
                if len(data) < 5:
                    return None
                height, width = struct.unpack(">HH", data[1:5])
                return width, height
            fh.seek(length - 2, 1)


def image_size(path: Path) -> tuple[int, int]:
    size = png_size(path) if path.suffix.lower() == ".png" else None
    if size is None and path.suffix.lower() in [".jpg", ".jpeg"]:
        size = jpeg_size(path)
    if size is None:
        raise ValueError(f"Could not determine image dimensions for {path}")
    return size


def add_picture_fit(slide, path: Path, left, top, width, height, *, mode: str = "contain"):
    img_w, img_h = image_size(path)
    img_aspect = img_w / img_h
    box_aspect = width / height

    if mode == "stretch":
        return slide.shapes.add_picture(str(path), left, top, width=width, height=height)

    if mode == "contain":
        if img_aspect >= box_aspect:
            new_w = width
            new_h = int(round(width / img_aspect))
        else:
            new_h = height
            new_w = int(round(height * img_aspect))
        new_left = left + int(round((width - new_w) / 2))
        new_top = top + int(round((height - new_h) / 2))
        return slide.shapes.add_picture(str(path), new_left, new_top, width=new_w, height=new_h)

    if mode == "cover":
        pic = slide.shapes.add_picture(str(path), left, top, width=width, height=height)
        if img_aspect > box_aspect:
            crop = (img_w - box_aspect * img_h) / (2 * img_w)
            pic.crop_left = crop
            pic.crop_right = crop
        elif img_aspect < box_aspect:
            crop = (img_h - img_w / box_aspect) / (2 * img_h)
            pic.crop_top = crop
            pic.crop_bottom = crop
        return pic

    raise ValueError(f"Unsupported picture fit mode: {mode}")


def replace_picture(slide, old_shape, path: Path, *, mode: str = "contain"):
    left, top, width, height = old_shape.left, old_shape.top, old_shape.width, old_shape.height
    delete_shape(old_shape)
    return add_picture_fit(slide, path, left, top, width, height, mode=mode)


def set_font_on_shape(shape, Pt, RGBColor, *, size=None, color=None, bold=None, italic=None) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if size is not None:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = RGBColor.from_string(color)
            if bold is not None:
                run.font.bold = bold
            if italic is not None:
                run.font.italic = italic


def set_font_on_paragraphs(
    shape,
    paragraph_indices: Sequence[int],
    Pt,
    RGBColor,
    *,
    size=None,
    color=None,
    bold=None,
    italic=None,
) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    wanted = set(paragraph_indices)
    for idx, paragraph in enumerate(shape.text_frame.paragraphs):
        if idx not in wanted:
            continue
        for run in paragraph.runs:
            if size is not None:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = RGBColor.from_string(color)
            if bold is not None:
                run.font.bold = bold
            if italic is not None:
                run.font.italic = italic


def add_caption_block(slide, pptx, left, top, width, height, name: str, role: str):
    MSO_SHAPE = pptx["MSO_SHAPE"]
    MSO_AUTO_SIZE = pptx["MSO_AUTO_SIZE"]
    PP_ALIGN = pptx["PP_ALIGN"]
    Pt = pptx["Pt"]
    RGBColor = pptx["RGBColor"]

    block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    block.fill.solid()
    block.fill.fore_color.rgb = RGBColor.from_string(ODI_BLUE_HEX)
    try:
        block.line.fill.background()
    except Exception:
        block.line.color.rgb = RGBColor.from_string(ODI_BLUE_HEX)

    tf = block.text_frame
    tf.clear()
    tf.margin_left = pptx["Inches"](0.08)
    tf.margin_right = pptx["Inches"](0.08)
    tf.margin_top = pptx["Inches"](0.04)
    tf.margin_bottom = pptx["Inches"](0.03)
    tf.word_wrap = True
    if MSO_AUTO_SIZE is not None:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    p_name = tf.paragraphs[0]
    p_name.alignment = PP_ALIGN.LEFT
    r_name = p_name.add_run()
    r_name.text = name
    r_name.font.name = "Arial"
    r_name.font.size = Pt(8.5)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(255, 255, 255)

    p_role = tf.add_paragraph()
    p_role.alignment = PP_ALIGN.LEFT
    r_role = p_role.add_run()
    r_role.text = role
    r_role.font.name = "Arial"
    r_role.font.size = Pt(6.4)
    r_role.font.color.rgb = RGBColor(255, 255, 255)
    return block


def add_plain_textbox(
    slide,
    pptx,
    text: str | Sequence[str],
    left,
    top,
    width,
    height,
    *,
    size: float = 12,
    color: str = ODI_BLUE_HEX,
    bold: bool | None = None,
    italic: bool | None = None,
):
    Pt = pptx["Pt"]
    RGBColor = pptx["RGBColor"]
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, line in enumerate(as_lines(text)):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        if bold is not None:
            run.font.bold = bold
        if italic is not None:
            run.font.italic = italic
    return box


def add_markdown_textbox(
    slide,
    pptx,
    paragraphs: Sequence[str],
    left,
    top,
    width,
    height,
    *,
    size: float = 12,
    color: str = ODI_BLUE_HEX,
    paragraph_space_after: float = 4,
    bold: bool = False,
):
    Pt = pptx["Pt"]
    RGBColor = pptx["RGBColor"]
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = pptx["Inches"](0.02)
    tf.margin_right = pptx["Inches"](0.02)
    tf.margin_top = pptx["Inches"](0.02)
    tf.margin_bottom = pptx["Inches"](0.02)

    for idx, paragraph_text in enumerate(paragraphs):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.space_after = Pt(paragraph_space_after)
        for segment, segment_bold in parse_bold_markers(paragraph_text):
            run = paragraph.add_run()
            run.text = segment
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor.from_string(color)
            run.font.bold = segment_bold if segment_bold is not None else bold
    return box


def add_agenda_item(slide, pptx, item: dict[str, str], left, top, width, height) -> None:
    Pt = pptx["Pt"]
    RGBColor = pptx["RGBColor"]
    Inches = pptx["Inches"]

    number_box = add_plain_textbox(
        slide,
        pptx,
        item["number"],
        left,
        top + Inches(0.02),
        Inches(0.42),
        height,
        size=15.5,
        bold=True,
    )
    number_box.text_frame.margin_left = 0
    number_box.text_frame.margin_right = 0

    title_box = slide.shapes.add_textbox(left + Inches(0.52), top, width - Inches(0.52), height)
    tf = title_box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    p_title = tf.paragraphs[0]
    r_title = p_title.add_run()
    r_title.text = item["title"]
    r_title.font.name = "Arial"
    r_title.font.size = Pt(15.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor.from_string(ODI_BLUE_HEX)

    p_subtitle = tf.add_paragraph()
    r_subtitle = p_subtitle.add_run()
    r_subtitle.text = f"{item['speaker']} · {item['duration']}"
    r_subtitle.font.name = "Arial"
    r_subtitle.font.size = Pt(11.3)
    r_subtitle.font.color.rgb = RGBColor.from_string(ODI_BLUE_HEX)


TEMPLATE_ARTIFACT_PATTERNS = [
    "Lorem ipsum",
    "Add name",
    "Example title",
    "Example subheading",
    "Click to add text",
    "Click to add title",
    "Add description",
    "Add map title here and paste map image below",
    "Add your SHORT title or statement here",
    "Replace this text with you section subheader",
    "Section title here",
    "Purus eget fringilla",
    "The outcomes and outputs",
    "Explanation of",
    "continued",
    "Watch the",
    "SOURCE:",
    "Stat or icon",
    "with a supporting statement here",
    "Add a short statement here",
]


EXACT_ARTIFACT_LABELS = ["Need", "Solution", "Impact"]


def clear_template_artifacts(slide) -> int:
    removed = delete_shapes_containing(slide, TEMPLATE_ARTIFACT_PATTERNS)
    # Remove leftover big faded section labels ("Need"/"Solution"/"Impact") that
    # only appear as exact single-word titles on the template content slides.
    for label in EXACT_ARTIFACT_LABELS:
        for shape in find_shapes_containing(slide, label, exact=True):
            delete_shape(shape)
            removed += 1
    return removed


def keep_and_reorder_slides(prs, pptx, warnings: list[str]) -> None:
    qn = pptx["qn"]
    sld_id_list = prs.slides._sldIdLst
    original_sld_ids = list(sld_id_list)

    if len(original_sld_ids) < max(TEMPLATE_INDICES):
        die(f"Template has only {len(original_sld_ids)} slides; need slide {max(TEMPLATE_INDICES)}.")
    if len(set(TEMPLATE_INDICES)) != len(TEMPLATE_INDICES):
        die("TEMPLATE_INDICES contains duplicates; this build expects 18 distinct template slides.")

    keep_zero_based = {idx - 1 for idx in TEMPLATE_INDICES}
    removed_rids = [
        sld_id.get(qn("r:id"))
        for zero_based, sld_id in enumerate(original_sld_ids)
        if zero_based not in keep_zero_based
    ]
    kept_sld_ids = [original_sld_ids[idx - 1] for idx in TEMPLATE_INDICES]

    for sld_id in original_sld_ids:
        sld_id_list.remove(sld_id)
    for sld_id in kept_sld_ids:
        sld_id_list.append(sld_id)

    for rid in removed_rids:
        if not rid:
            continue
        try:
            prs.part.drop_rel(rid)
        except Exception as exc:
            warnings.append(f"Could not drop relationship {rid} for a removed template slide: {exc}")


def build_slide_01(slide, pptx, warnings):
    content = CONTENT["s01"]
    replace_first(slide, "This is the title of the presentation", content["title"], warnings, "slide 01")
    replace_first(
        slide,
        "[name of presenter + contact details]",
        [
            content["strap"],
            content["meta"][0],
            content["meta"][1],
            content["meta"][2],
        ],
        warnings,
        "slide 01",
    )
    for shape in find_shapes_containing(slide, content["strap"]):
        set_font_on_paragraphs(shape, [0], pptx["Pt"], pptx["RGBColor"], italic=True)


def build_slide_02(slide, pptx, warnings):
    content = CONTENT["s02"]
    Inches = pptx["Inches"]

    title_candidates = []
    for shape in walk_shapes(slide.shapes):
        text = shape_text(shape)
        normalized = normalize_text(text)
        if text and (normalized in {"contents", "run of show", "show"} or "contents" in normalized):
            title_candidates.append(shape)

    if title_candidates:
        title_candidates = sorted(title_candidates, key=lambda shape: (shape.top, shape.left))
        title_shape = title_candidates[0]
        set_text(title_shape, content["title"])
        title_shape.left = Inches(0.34)
        title_shape.top = Inches(0.34)
        title_shape.width = Inches(7.5)
        title_shape.height = Inches(0.52)
        set_font_on_shape(title_shape, pptx["Pt"], pptx["RGBColor"], size=22, color=ODI_BLUE_HEX, bold=True)
        for extra in title_candidates[1:]:
            delete_shape(extra)
    else:
        add_plain_textbox(
            slide,
            pptx,
            content["title"],
            Inches(0.34),
            Inches(0.34),
            Inches(7.5),
            Inches(0.52),
            size=22,
            bold=True,
        )

    delete_shapes_containing(slide, ["Example title", "Example subheading"])
    agenda_top = Inches(1.24)
    row_gap = Inches(0.90)
    for idx, item in enumerate(content["agenda"]):
        add_agenda_item(
            slide,
            pptx,
            item,
            Inches(1.08),
            agenda_top + row_gap * idx,
            Inches(11.15),
            Inches(0.64),
        )

    add_plain_textbox(
        slide,
        pptx,
        content["footer"],
        Inches(0.45),
        Inches(6.50),
        Inches(10.9),
        Inches(0.35),
        size=8.5,
        color=ODI_BLUE_HEX,
    )


def build_slide_03(slide, pptx, warnings):
    content = CONTENT["s03"]
    replace_first(slide, "ODI team", content["kicker"], warnings, "slide 03")
    delete_shapes_containing(slide, ["Add name", "Watch the"])
    for pic in list(picture_shapes(slide, pptx["MSO_SHAPE_TYPE"])):
        delete_shape(pic)

    Inches = pptx["Inches"]
    card_w = Inches(3.75)
    img_w = Inches(2.05)
    img_h = Inches(2.30)
    cap_w = Inches(3.45)
    cap_h = Inches(0.58)
    col_lefts = [Inches(0.35), Inches(4.80), Inches(9.05)]
    row_tops = [Inches(1.04), Inches(3.98)]

    for idx, speaker in enumerate(content["speakers"]):
        col = idx % 3
        row = idx // 3
        image_left = col_lefts[col] + int(round((card_w - img_w) / 2))
        image_top = row_tops[row]
        cap_left = col_lefts[col] + int(round((card_w - cap_w) / 2))
        cap_top = image_top + img_h + Inches(0.06)
        add_picture_fit(slide, BW_DIR / speaker["image"], image_left, image_top, img_w, img_h, mode="cover")
        add_caption_block(slide, pptx, cap_left, cap_top, cap_w, cap_h, speaker["name"], speaker["role"])


def build_slide_04(slide, pptx, warnings):
    content = CONTENT["s04"]
    Inches = pptx["Inches"]
    MSO_SHAPE = pptx["MSO_SHAPE"]
    RGBColor = pptx["RGBColor"]

    delete_shapes_containing(slide, ["ODI partners and clients", "The ODI — steward", "Mission-driven non-profit"])

    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.34), Inches(2.08))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string(ODI_BLUE_HEX)
    try:
        band.line.fill.background()
    except Exception:
        band.line.color.rgb = RGBColor.from_string(ODI_BLUE_HEX)

    add_plain_textbox(
        slide,
        pptx,
        content["kicker"],
        Inches(0.36),
        Inches(0.24),
        Inches(7.4),
        Inches(0.24),
        size=9.0,
        color="FFFFFF",
    )
    add_plain_textbox(
        slide,
        pptx,
        content["title"],
        Inches(0.36),
        Inches(0.55),
        Inches(8.0),
        Inches(0.45),
        size=25.0,
        color="FFFFFF",
        bold=True,
    )
    add_plain_textbox(
        slide,
        pptx,
        content["body"],
        Inches(0.36),
        Inches(1.10),
        Inches(11.7),
        Inches(0.70),
        size=9.8,
        color="FFFFFF",
    )
    add_plain_textbox(
        slide,
        pptx,
        content["trusted_label"],
        Inches(0.36),
        Inches(2.20),
        Inches(3.8),
        Inches(0.28),
        size=11.0,
        color="FFFFFF",
        bold=True,
    )


def build_slide_05(slide, pptx, warnings):
    delete_shapes_containing(slide, ["Add description", "SOURCE:", "Watch the"])
    pics = picture_shapes(slide, pptx["MSO_SHAPE_TYPE"])
    for pic in pics:
        delete_shape(pic)
    add_picture_fit(
        slide,
        IMG_DIR / "data-spectrum.jpg",
        pptx["Inches"](0.0),
        pptx["Inches"](0.0),
        pptx["Inches"](13.34),
        pptx["Inches"](7.50),
        mode="contain",
    )


def build_slide_06(slide, pptx, warnings):
    content = CONTENT["s06"]
    replace_first(slide, "Section title", content["title"], warnings, "slide 06")
    replace_first(
        slide,
        "Replace this text with you section subheader",
        [content["subtitle"], content["who"]],
        warnings,
        "slide 06",
    )


def build_slide_07(slide, pptx, warnings):
    content = CONTENT["s07"]
    headings = sorted_text_shapes(slide, "Column heading", exact=True)
    if headings:
        set_text(headings[0], content["title"])
        for extra in headings[1:]:
            delete_shape(extra)
    else:
        warnings.append("slide 07: could not find column heading placeholders")

    bodies = sorted_text_shapes(slide, "Two column layout with heading")
    if bodies:
        set_markdown_text(
            bodies[0],
            content["body"],
        )
        for extra in bodies[1:]:
            delete_shape(extra)
    else:
        warnings.append("slide 07: could not find body placeholder")

    delete_shapes_containing(slide, ["Watch the"])
    pics = picture_shapes(slide, pptx["MSO_SHAPE_TYPE"])
    if pics:
        replace_picture(slide, pics[-1], IMG_DIR / "solid-pods.png", mode="contain")
    else:
        warnings.append("slide 07: could not find example picture; adding Solid pods image in fallback box")
        add_picture_fit(
            slide,
            IMG_DIR / "solid-pods.png",
            pptx["Inches"](8.20),
            pptx["Inches"](0.60),
            pptx["Inches"](4.75),
            pptx["Inches"](5.90),
            mode="contain",
        )
    add_plain_textbox(
        slide,
        pptx,
        content["figure_caption"],
        pptx["Inches"](8.15),
        pptx["Inches"](6.08),
        pptx["Inches"](4.55),
        pptx["Inches"](0.28),
        size=8.7,
        color=ODI_BLUE_HEX,
        bold=True,
    )


def build_slide_08(slide, pptx, warnings):
    content = CONTENT["s08"]
    Inches = pptx["Inches"]
    replace_first(
        slide,
        "Heading for key points slide",
        content["title"],
        warnings,
        "slide 08",
    )
    delete_shapes_containing(slide, ["Purus eget fringilla"])
    add_markdown_textbox(
        slide,
        pptx,
        [content["lead"], *[f"• {bullet}" for bullet in content["bullets"]], content["fix"], content["ai_line"]],
        Inches(0.72),
        Inches(1.18),
        Inches(6.60),
        Inches(5.25),
        size=12.5,
        paragraph_space_after=3.5,
    )
    add_picture_fit(
        slide,
        IMG_DIR / "walled-gardens.png",
        Inches(7.80),
        Inches(1.45),
        Inches(4.70),
        Inches(3.70),
        mode="contain",
    )
    add_plain_textbox(
        slide,
        pptx,
        content["figure_caption"],
        Inches(7.80),
        Inches(5.20),
        Inches(4.70),
        Inches(0.30),
        size=8.0,
        color=ODI_BLUE_HEX,
        italic=True,
    )


def build_slide_09(slide, pptx, warnings):
    content = CONTENT["s09"]
    replace_first(slide, "Add map title here and paste map image below", content["title"], warnings, "slide 09")
    delete_shapes_containing(slide, ["Watch the"])
    add_plain_textbox(
        slide,
        pptx,
        content["flag"],
        pptx["Inches"](0.70),
        pptx["Inches"](0.72),
        pptx["Inches"](1.35),
        pptx["Inches"](0.26),
        size=10.0,
        color=ODI_BLUE_HEX,
        bold=True,
    )
    pic = add_picture_fit(
        slide,
        IMG_DIR / "solid-demo.jpg",
        pptx["Inches"](0.70),
        pptx["Inches"](1.05),
        pptx["Inches"](11.95),
        pptx["Inches"](5.45),
        mode="contain",
    )
    try:
        pic.click_action.hyperlink.address = FIGMA_URL
    except Exception as exc:
        warnings.append(f"slide 09: could not attach Figma hyperlink to image: {exc}")
    add_plain_textbox(
        slide,
        pptx,
        content["footer"],
        pptx["Inches"](0.70),
        pptx["Inches"](6.55),
        pptx["Inches"](7.0),
        pptx["Inches"](0.32),
        size=9.0,
        color=ODI_BLUE_HEX,
    )


def build_slide_10(slide, pptx, warnings):
    content = CONTENT["s10"]
    replace_first(slide, "Founders", content["kicker"], warnings, "slide 10")
    replace_first(
        slide,
        "The emergence of the latest generative AI",
        [content["title"], content["footer"]],
        warnings,
        "slide 10",
    )

    pics = picture_shapes(slide, pptx["MSO_SHAPE_TYPE"])
    if len(pics) >= 2:
        replace_picture(slide, pics[0], BW_DIR / content["people"][0]["image"], mode="cover")
        replace_picture(slide, pics[1], BW_DIR / content["people"][1]["image"], mode="cover")
    else:
        warnings.append(f"slide 10: expected 2 founder pictures, found {len(pics)}")

    captions = sorted(
        [shape for shape in walk_shapes(slide.shapes) if "Executive Chair and Co-founder" in shape_text(shape) or "President and Co-founder" in shape_text(shape)],
        key=lambda shape: shape.left,
    )
    values = [[person["name"], person["role"]] for person in content["people"]]
    if len(captions) >= 2:
        for shape, lines in zip(captions[:2], values):
            set_text(shape, lines)
            set_font_on_shape(shape, pptx["Pt"], pptx["RGBColor"], size=7.4)
    else:
        warnings.append(f"slide 10: expected 2 founder caption shapes, found {len(captions)}")


def build_slide_11(slide, pptx, warnings):
    content = CONTENT["s11"]
    replace_first(slide, "Add your SHORT title or statement here", content["title"], warnings, "slide 11")
    title_matches = find_shapes_containing(slide, content["title"])
    if title_matches:
        title_shape = title_matches[0]
        title_shape.top = pptx["Inches"](1.65)
        title_shape.height = pptx["Inches"](1.05)
        set_font_on_shape(title_shape, pptx["Pt"], pptx["RGBColor"], size=40, color=ODI_BLUE_HEX, bold=True)

    add_plain_textbox(
        slide,
        pptx,
        content["kicker"],
        pptx["Inches"](1.18),
        pptx["Inches"](0.76),
        pptx["Inches"](10.75),
        pptx["Inches"](0.36),
        size=13.0,
        color=ODI_BLUE_HEX,
        bold=True,
    )
    add_markdown_textbox(
        slide,
        pptx,
        content["body"],
        pptx["Inches"](1.18),
        pptx["Inches"](3.52),
        pptx["Inches"](10.75),
        pptx["Inches"](1.10),
        size=17.0,
        color=ODI_BLUE_HEX,
    )
    add_plain_textbox(
        slide,
        pptx,
        content["footer"],
        pptx["Inches"](1.18),
        pptx["Inches"](5.90),
        pptx["Inches"](10.75),
        pptx["Inches"](0.35),
        size=11.0,
        color=ODI_BLUE_HEX,
    )


def build_slide_12(slide, pptx, warnings):
    content = CONTENT["s12"]
    replace_first(slide, "Section title here", content["title"], warnings, "slide 12")
    add_plain_textbox(
        slide,
        pptx,
        [content["subtitle"], content["who"]],
        pptx["Inches"](1.20),
        pptx["Inches"](4.92),
        pptx["Inches"](10.55),
        pptx["Inches"](0.70),
        size=15.0,
        color=ODI_BLUE_HEX,
    )


def build_slide_13(slide, pptx, warnings):
    content = CONTENT["s13"]
    replace_first(slide, "Click to add title", content["title"], warnings, "slide 13")
    title_matches = find_shapes_containing(slide, content["title"])
    if title_matches:
        title_shape = title_matches[0]
        title_shape.left = pptx["Inches"](0.34)
        title_shape.top = pptx["Inches"](0.58)
        title_shape.width = pptx["Inches"](11.1)
        title_shape.height = pptx["Inches"](0.55)
        set_font_on_shape(title_shape, pptx["Pt"], pptx["RGBColor"], size=28, color=ODI_BLUE_HEX, bold=True)
    add_plain_textbox(
        slide,
        pptx,
        content["kicker"],
        pptx["Inches"](0.34),
        pptx["Inches"](0.28),
        pptx["Inches"](6.6),
        pptx["Inches"](0.26),
        size=10.0,
        color=ODI_BLUE_HEX,
        bold=True,
    )
    add_markdown_textbox(
        slide,
        pptx,
        [content["lead"]],
        pptx["Inches"](0.55),
        pptx["Inches"](1.18),
        pptx["Inches"](11.85),
        pptx["Inches"](0.72),
        size=10.4,
        paragraph_space_after=0,
    )

    icon_shapes = sorted_text_shapes(slide, "Stat or icon", exact=True)
    for shape, card in zip(icon_shapes, content["cards"]):
        set_text(shape, card["label"])
    if len(icon_shapes) != 3:
        warnings.append(f"slide 13: expected 3 stat/icon circles, found {len(icon_shapes)}")

    caption_shapes = sorted_text_shapes(slide, "Click to add text", exact=True)
    for shape, card in zip(caption_shapes, content["cards"]):
        set_text(shape, [card["title"], card["body"]])
        set_font_on_paragraphs(shape, [0], pptx["Pt"], pptx["RGBColor"], size=10.0, bold=True)
        set_font_on_paragraphs(shape, [1], pptx["Pt"], pptx["RGBColor"], size=8.0)
    if len(caption_shapes) != 3:
        warnings.append(f"slide 13: expected 3 caption placeholders, found {len(caption_shapes)}")


def build_focus_slide(slide, pptx, warnings, *, slide_num: str, content: dict):
    Inches = pptx["Inches"]

    replace_first_of(
        slide,
        ["The need", "The solution", "The impact"],
        content["kicker"],
        warnings,
        slide_num,
        "focus kicker",
    )
    replace_first_of(
        slide,
        ["Need", "Solution", "Impact"],
        content["title"],
        warnings,
        slide_num,
        "focus title",
    )

    kicker_matches = find_shapes_containing(slide, content["kicker"])
    if kicker_matches:
        kicker_shape = kicker_matches[0]
        kicker_shape.left = Inches(0.34)
        kicker_shape.top = Inches(0.30)
        kicker_shape.width = Inches(7.0)
        kicker_shape.height = Inches(0.25)
        set_font_on_shape(kicker_shape, pptx["Pt"], pptx["RGBColor"], size=10, color=ODI_BLUE_HEX, bold=True)

    title_matches = find_shapes_containing(slide, content["title"])
    if title_matches:
        title_shape = title_matches[0]
        title_shape.left = Inches(0.34)
        title_shape.top = Inches(0.70)
        title_shape.width = Inches(11.3)
        title_shape.height = Inches(0.62)
        set_font_on_shape(title_shape, pptx["Pt"], pptx["RGBColor"], size=26, color=ODI_BLUE_HEX, bold=True)
        for extra in title_matches[1:]:
            delete_shape(extra)

    delete_shapes_containing(
        slide,
        [
            "Lorem ipsum",
            "The outcomes and outputs",
            "Explanation of",
            "continued",
            "The need",
            "The solution",
            "The impact",
        ],
    )

    paragraphs = [f"**{point['heading']}** {point['body']}" for point in content["points"]]
    add_markdown_textbox(
        slide,
        pptx,
        paragraphs,
        Inches(0.90),
        Inches(1.72),
        Inches(10.85),
        Inches(4.40),
        size=15.0 if slide_num == "slide 16" else 16.5,
        paragraph_space_after=12,
    )


def build_slide_14(slide, pptx, warnings):
    build_focus_slide(slide, pptx, warnings, slide_num="slide 14", content=CONTENT["s14"])


def build_slide_15(slide, pptx, warnings):
    build_focus_slide(slide, pptx, warnings, slide_num="slide 15", content=CONTENT["s15"])


def build_slide_16(slide, pptx, warnings):
    build_focus_slide(slide, pptx, warnings, slide_num="slide 16", content=CONTENT["s16"])


def build_slide_17(slide, pptx, warnings):
    content = CONTENT["s17"]
    replace_first(
        slide,
        "Add a short statement here",
        content["title"],
        warnings,
        "slide 17",
    )
    replace_first(
        slide,
        "with a supporting statement here",
        [
            content["badge"],
            content["register"],
            content["url"],
            content["body"],
            content["cochairs"],
        ],
        warnings,
        "slide 17",
    )
    title_matches = find_shapes_containing(slide, content["title"])
    if title_matches:
        title_shape = title_matches[0]
        title_shape.top = pptx["Inches"](1.00)
        title_shape.height = pptx["Inches"](0.92)
        set_font_on_shape(title_shape, pptx["Pt"], pptx["RGBColor"], size=31, color=ODI_BLUE_HEX, bold=True)
    body_matches = find_shapes_containing(slide, content["body"])
    if body_matches:
        body_shape = body_matches[0]
        body_shape.top = pptx["Inches"](2.35)
        body_shape.height = pptx["Inches"](3.30)
        set_font_on_paragraphs(body_shape, [0], pptx["Pt"], pptx["RGBColor"], size=11.0, color=ODI_BLUE_HEX, bold=True)
        set_font_on_paragraphs(body_shape, [1], pptx["Pt"], pptx["RGBColor"], size=18, color=ODI_BLUE_HEX, bold=True)
        set_font_on_paragraphs(body_shape, [2], pptx["Pt"], pptx["RGBColor"], size=22, color=ODI_BLUE_HEX, bold=True)
        set_font_on_paragraphs(body_shape, [3], pptx["Pt"], pptx["RGBColor"], size=13.0, color=ODI_BLUE_HEX)
        set_font_on_paragraphs(body_shape, [4], pptx["Pt"], pptx["RGBColor"], size=10.2, color=ODI_BLUE_HEX)


def build_slide_18(slide, pptx, warnings):
    content = CONTENT["s18"]
    replace_first(slide, "Section title", content["title_lines"], warnings, "slide 18")
    replace_first(
        slide,
        "Replace this text with you section subheader",
        content["subheader"],
        warnings,
        "slide 18",
    )
    title_matches = find_shapes_containing(slide, content["title_lines"][0])
    if title_matches:
        title_shape = title_matches[0]
        title_shape.top = pptx["Inches"](0.58)
        title_shape.height = pptx["Inches"](2.10)
        set_font_on_shape(title_shape, pptx["Pt"], pptx["RGBColor"], size=42, color="FFFFFF", bold=True)
    subheader_matches = find_shapes_containing(slide, content["subheader"])
    if subheader_matches:
        subheader_shape = subheader_matches[0]
        subheader_shape.top = pptx["Inches"](4.45)
        subheader_shape.height = pptx["Inches"](1.25)
        set_font_on_shape(subheader_shape, pptx["Pt"], pptx["RGBColor"], size=16.0, color="FFFFFF")


SLIDE_BUILDERS = [
    ("Title", 3, build_slide_01),
    ("Agenda", 4, build_slide_02),
    ("Meet the speakers", 59, build_slide_03),
    ("ODI steward of Solid", 46, build_slide_04),
    ("The Data Spectrum", 26, build_slide_05),
    ("Divider: Introduction to Solid", 15, build_slide_06),
    ("What is Solid", 32, build_slide_07),
    ("The problem", 9, build_slide_08),
    ("Live demo", 53, build_slide_09),
    ("Fireside", 57, build_slide_10),
    ("Q&A", 18, build_slide_11),
    ("Divider: Community of Practice", 16, build_slide_12),
    ("What we're building", 41, build_slide_13),
    ("Focus A: pilots", 6, build_slide_14),
    ("Focus B: standards", 7, build_slide_15),
    ("Focus C: certification", 8, build_slide_16),
    ("Get involved", 20, build_slide_17),
    ("Thank you", 14, build_slide_18),
]


def main() -> int:
    print("Building ODI webinar deck")
    print(f"Template: {TEMPLATE_PATH}")
    print(f"Output:   {OUTPUT_PATH}")

    validate_inputs()
    pptx = import_pptx()
    Presentation = pptx["Presentation"]

    warnings: list[str] = []
    prs = Presentation(str(TEMPLATE_PATH))
    print(f"Template opened: {len(prs.slides)} slides, {prs.slide_width}×{prs.slide_height} EMU")

    keep_and_reorder_slides(prs, pptx, warnings)
    print(f"Kept/reordered {len(prs.slides)} slides: {TEMPLATE_INDICES}")

    if len(prs.slides) != len(SLIDE_BUILDERS):
        die(f"After pruning, deck has {len(prs.slides)} slides; expected {len(SLIDE_BUILDERS)}.")

    for slide_number, (label, template_index, builder) in enumerate(SLIDE_BUILDERS, start=1):
        print(f"[slide {slide_number:02d}] {label} <- template slide {template_index}")
        builder(prs.slides[slide_number - 1], pptx, warnings)
        clear_template_artifacts(prs.slides[slide_number - 1])

    prs.save(str(OUTPUT_PATH))
    print(f"Saved: {OUTPUT_PATH}")

    if warnings:
        print("\nWarnings:")
        for warning in warnings:
            print(f" - {warning}")
    else:
        print("Warnings: none")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as exc:
        print(f"\nERROR: build failed: {exc}", file=sys.stderr)
        print("Traceback follows so the next pass can fix the exact failing shape/API call:", file=sys.stderr)
        traceback.print_exc()
        raise
